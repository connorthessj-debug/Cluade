"""End-to-end smoke test: the generator produces a valid xlsx with 3 sheets."""
import os

import openpyxl
import pytest

import generate_schedule as gs


def test_create_option_sheet_smoke(tmp_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    gs.create_option_sheet(wb, "Option A - Wed", "A")

    assert "Option A - Wed" in wb.sheetnames
    ws = wb["Option A - Wed"]
    # Title in row 1
    assert "Option A" in str(ws.cell(row=1, column=1).value)

    # Round-trip save/load
    out = tmp_path / "out.xlsx"
    wb.save(out)
    reopened = openpyxl.load_workbook(out)
    assert "Option A - Wed" in reopened.sheetnames


def test_main_creates_three_option_sheets(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    gs.main()
    out = tmp_path / "schedule_options.xlsx"
    assert out.exists()
    wb = openpyxl.load_workbook(out)
    assert set(wb.sheetnames) == {
        "Option A - Wed",
        "Option B - Mon",
        "Option C - Tue",
    }


def test_write_section_advances_row(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    start = 3
    end = gs.write_section(ws, start, gs.get_afternoon_shift())
    # Must advance at least past the title + header + staff + summary rows
    assert end > start + 6
